"""
Excel Exporter - خروجی اکسل از داده‌های پارکینگ
"""

import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from utils import to_shamsi

class ExcelExporter:
    """تولید فایل Excel از گزارش‌ها"""

    def __init__(self, db):
        self.db = db
        self.export_dir = "exports"
        os.makedirs(self.export_dir, exist_ok=True)

    def export_history(self, date_from=None, date_to=None, filename=None):
        """خروجی Excel از تاریخچه"""
        wb = Workbook()
        ws = wb.active
        ws.title = "تاریخچه پارکینگ"
        ws.sheet_view.rightToLeft = True

        # ===== هدر =====
        headers = [
            'ردیف', 'پلاک', 'کارت', 'نوع', 'استان',
            'زمان ورود', 'زمان خروج', 'مدت (ساعت)',
            'هزینه پایه', 'تخفیف', 'هزینه نهایی',
            'اپراتور', 'روش پرداخت'
        ]

        header_font = Font(name='Tahoma', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='2c3e50', end_color='2c3e50', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

        # ===== داده‌ها =====
        history = self.db.get_history(
            date_from=date_from,
            date_to=date_to,
            per_page=10000
        )

        for i, record in enumerate(history['records'], 1):
            row_data = [
                i,
                record['plate_number'],
                record.get('card_number', '-'),
                record.get('plate_type', 'شخصی'),
                record.get('province', ''),
                to_shamsi(record['entry_time'], '%Y/%m/%d %H:%M'),
                to_shamsi(record['exit_time'], '%Y/%m/%d %H:%M'),
                f"{record['duration_hours']:.1f}",
                f"{record['cost']:,.0f}",
                f"{record['discount_percent']}%" if record['discount_percent'] > 0 else "-",
                f"{record['final_cost']:,.0f}",
                record.get('operator_name', '-'),
                record.get('payment_method', 'نقدی')
            ]

            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=i + 1, column=col, value=value)
                cell.font = Font(name='Tahoma', size=10)
                cell.alignment = Alignment(horizontal='center', vertical='center')

                # رنگ‌آمیزی هزینه نهایی
                if col == 11 and record['final_cost'] > 0:
                    cell.font = Font(name='Tahoma', size=10, bold=True, color='c0392b')

        # ===== تنظیم عرض ستون‌ها =====
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 15

        # ===== ذخیره =====
        if not filename:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f"parking_report_{timestamp}.xlsx"

        filepath = os.path.join(self.export_dir, filename)
        wb.save(filepath)

        return filepath

    def export_daily_summary(self, date=None, filename=None):
        """خروجی Excel از خلاصه روزانه"""
        if not date:
            date = datetime.now().strftime('%Y-%m-%d')

        wb = Workbook()
        ws = wb.active
        ws.title = "خلاصه روزانه"
        ws.sheet_view.rightToLeft = True

        # ===== هدر =====
        ws.merge_cells('A1:D1')
        title_cell = ws['A1']
        title_cell.value = f"گزارش روزانه پارکینگ - {date}"
        title_cell.font = Font(name='Tahoma', size=14, bold=True, color='2c3e50')
        title_cell.alignment = Alignment(horizontal='center', vertical='center')

        # ===== آمار =====
        stats = self.db.get_statistics(date)

        ws['A3'] = "شاخص"
        ws['B3'] = "مقدار"
        ws['A3'].font = Font(bold=True)
        ws['B3'].font = Font(bold=True)

        rows = [
            ("تعداد خودروهای خارج‌شده", stats['daily']['count']),
            ("درآمد کل (تومان)", f"{stats['daily']['income']:,.0f}"),
            ("میانگین مدت توقف (ساعت)", f"{stats['daily']['avg_duration']:.1f}"),
            ("تعداد خودروهای منحصربه‌فرد", stats['daily']['unique_cars']),
            ("خودروهای حاضر", stats['active_cars']),
        ]

        for i, (label, value) in enumerate(rows, 4):
            ws.cell(row=i, column=1, value=label).font = Font(name='Tahoma', size=11)
            ws.cell(row=i, column=2, value=value).font = Font(name='Tahoma', size=11, bold=True)

        # ===== تنظیم عرض =====
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 20

        # ===== ذخیره =====
        if not filename:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f"daily_summary_{date}_{timestamp}.xlsx"

        filepath = os.path.join(self.export_dir, filename)
        wb.save(filepath)

        return filepath